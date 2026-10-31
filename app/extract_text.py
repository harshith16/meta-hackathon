"""Text extraction from uploaded documents and images. Pure-Python, Windows-safe.

PDF via pypdf (with scanned-page image OCR fallback), DOCX via python-docx,
TXT/MD/CSV/JSON via plain decode, and images (JPG, PNG, WEBP) via Vision LLM.
"""

import io
import os


def _mime_for_filename(filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    return {
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".tiff": "image/tiff",
    }.get(ext, "image/jpeg")


def _extract_txt(data: bytes, filename: str = "") -> str | None:
    text = data.decode("utf-8", errors="replace").strip()
    return text or None


def _extract_pdf(data: bytes, filename: str = "") -> str | None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            text = page.extract_text()
        except Exception:
            text = None
        if text:
            parts.append(text)
    joined = "\n".join(parts).strip()
    if joined:
        return joined

    # If digital text was empty (scanned PDF), extract embedded page images
    try:
        from app.llm import transcribe_image

        image_parts = []
        for page in reader.pages:
            for img in getattr(page, "images", []):
                transcribed = transcribe_image(img.data, "image/jpeg")
                if transcribed:
                    image_parts.append(transcribed)
        if image_parts:
            return "\n\n".join(image_parts).strip()
    except Exception:
        pass

    return None


def _extract_docx(data: bytes, filename: str = "") -> str | None:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    joined = "\n".join(parts).strip()
    return joined or None


def _extract_image(data: bytes, filename: str = "") -> str | None:
    try:
        from app.llm import transcribe_image

        mime = _mime_for_filename(filename)
        return transcribe_image(data, mime)
    except Exception:
        return None


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".text": _extract_txt,
    ".md": _extract_txt,
    ".markdown": _extract_txt,
    ".csv": _extract_txt,
    ".json": _extract_txt,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".jpg": _extract_image,
    ".jpeg": _extract_image,
    ".png": _extract_image,
    ".webp": _extract_image,
    ".bmp": _extract_image,
    ".tiff": _extract_image,
}


def extract_text(filename, data: bytes) -> str | None:
    """Extract text from document or image bytes, or None when not possible."""
    if not data:
        return None
    ext = os.path.splitext(filename or "")[1].lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        return None
    try:
        return fn(data, filename or "")
    except Exception:
        return None
